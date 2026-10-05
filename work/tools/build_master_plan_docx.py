from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


RTL_FONT = "Arial"
CODE_FONT = "Consolas"
TEAL = "0F5B57"
PALE = "F3F7F6"
CREAM = "FAF6EE"
GRID = "D9D9D9"


def set_run_font(run, name=RTL_FONT, size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_rtl(paragraph, *, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_ltr(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is not None:
        ppr.remove(bidi)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def clean_heading(text):
    text = re.sub(r"[`*_]+", "", text)
    text = text.replace("—", " ").replace("–", " ")
    text = text.replace("«", "").replace("»", "")
    return re.sub(r"\s+", " ", text).strip(" .:-")


INLINE_RE = re.compile(r"(\*\*.*?\*\*|`.*?`)")


def add_inline(paragraph, text, *, default_size=11.5, default_color="000000"):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, bold=True, color=default_color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=CODE_FONT, size=max(default_size - 1, 9), color=default_color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=default_size, color=default_color)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = RTL_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), RTL_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), RTL_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), RTL_FONT)
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    style_specs = {
        "Title": (24, 14, 10),
        "Subtitle": (12, 5, 5),
        "Heading 1": (18, 14, 8),
        "Heading 2": (14, 11, 5),
        "Heading 3": (12.5, 8, 4),
    }
    for name, (size, before, after) in style_specs.items():
        style = styles[name]
        style.font.name = RTL_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), RTL_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), RTL_FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), RTL_FONT)
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

        style_ppr = style._element.get_or_add_pPr()
        border = style_ppr.find(qn("w:pBdr"))
        if border is not None:
            style_ppr.remove(border)

    styles["Heading 1"].paragraph_format.page_break_before = True

    settings = doc.settings._element
    theme_font_lang = settings.find(qn("w:themeFontLang"))
    if theme_font_lang is None:
        theme_font_lang = OxmlElement("w:themeFontLang")
        settings.append(theme_font_lang)
    theme_font_lang.set(qn("w:val"), "ar-SA")
    theme_font_lang.set(qn("w:bidi"), "ar-SA")


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    set_rtl(title, align=WD_ALIGN_PARAGRAPH.CENTER)
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    run = title.add_run("المرجع التنفيذي الشامل لتطبيق مدرب المقابلات القيادية")
    set_run_font(run, size=24, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    set_rtl(subtitle, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = subtitle.add_run("خطة المراحل الثلاث وملف التسليم للوكيل المنفذ")
    set_run_font(run, size=13)

    doc.add_paragraph()
    metadata = [
        "إصدار الوثيقة 1.1",
        "الإصدار الحالي للتطبيق 0.5.0 المرحلة الثانية التجريبية",
        "4 أكتوبر 2026",
    ]
    for item in metadata:
        p = doc.add_paragraph()
        set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, item, default_size=11)

    doc.add_page_break()


def add_body_paragraph(doc, text, *, style=None, indent=0):
    p = doc.add_paragraph(style=style)
    set_rtl(p)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.right_indent = Inches(indent)
    p.paragraph_format.widow_control = True
    add_inline(p, text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.right_indent = Inches(0.22 + 0.22 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    bullet = p.add_run("• ")
    set_run_font(bullet, size=11.5, bold=True, color=TEAL)
    add_inline(p, text)


def add_numbered(doc, number, text, level=0):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.right_indent = Inches(0.25 + 0.22 * level)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    lead = p.add_run(f"{number}. ")
    set_run_font(lead, size=11.5, bold=True, color=TEAL)
    add_inline(p, text)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    contains_arabic = any(re.search(r"[\u0600-\u06FF]", line) for line in lines)
    if contains_arabic:
        set_rtl(p)
    else:
        set_ltr(p)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name=RTL_FONT if contains_arabic else CODE_FONT, size=9.5, color="222222")
        if index < len(lines) - 1:
            run.add_break()
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F4F4")
    ppr.append(shd)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_rtl(table)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    for r_index, source_row in enumerate(rows):
        row = table.rows[r_index]
        for c_index in range(cols):
            text = source_row[c_index] if c_index < len(source_row) else ""
            cell = row.cells[c_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_index == 0:
                set_cell_shading(cell, TEAL)
            elif r_index % 2 == 0:
                set_cell_shading(cell, PALE)
            else:
                set_cell_shading(cell, "FFFFFF")
            p = cell.paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_after = Pt(0)
            color = "FFFFFF" if r_index == 0 else "000000"
            add_inline(p, text, default_size=10.2, default_color=color)
            if r_index == 0:
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", part.replace(" ", "")) for part in parts):
            rows.append(parts)
        index += 1
    return rows, index


def build(markdown_path: Path, output_path: Path):
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    if lines and lines[0].startswith("# "):
        try:
            frontmatter_end = lines.index("---")
            lines = lines[frontmatter_end + 1:]
        except ValueError:
            pass
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    in_code = False
    code_lines = []
    index = 0
    first_h1 = True
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            set_rtl(p)
            add_inline(p, clean_heading(stripped[4:]), default_size=12.5)
            index += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            set_rtl(p)
            add_inline(p, clean_heading(stripped[3:]), default_size=14)
            index += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            if first_h1:
                p.paragraph_format.page_break_before = False
                first_h1 = False
            set_rtl(p)
            add_inline(p, clean_heading(stripped[2:]), default_size=18)
            index += 1
            continue
        if stripped.startswith("> "):
            p = add_body_paragraph(doc, stripped[2:])
            p.paragraph_format.right_indent = Inches(0.35)
            p.paragraph_format.left_indent = Inches(0.35)
            for run in p.runs:
                run.italic = True
            index += 1
            continue
        bullet_match = re.match(r"^\s*-\s+(.*)$", raw)
        if bullet_match:
            leading = len(raw) - len(raw.lstrip(" "))
            add_bullet(doc, bullet_match.group(1), level=max(0, leading // 2))
            index += 1
            continue
        number_match = re.match(r"^\s*(\d+)\.\s+(.*)$", raw)
        if number_match:
            leading = len(raw) - len(raw.lstrip(" "))
            add_numbered(doc, number_match.group(1), number_match.group(2), level=max(0, leading // 2))
            index += 1
            continue

        add_body_paragraph(doc, stripped)
        index += 1

    if code_lines:
        add_code_block(doc, code_lines)

    core = doc.core_properties
    core.title = "المرجع التنفيذي الشامل لتطبيق مدرب المقابلات القيادية"
    core.subject = "خطة المراحل الثلاث وملف التسليم للوكيل المنفذ"
    core.author = ""
    core.keywords = "مقابلات قيادية، خطة تنفيذ، تطبيق، ذكاء اصطناعي"
    core.comments = ""
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_master_plan_docx.py INPUT.md OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
